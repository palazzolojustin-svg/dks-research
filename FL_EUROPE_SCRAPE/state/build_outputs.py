import json, os, re
root = r"C:\Users\palaz\Downloads\DKS_RESEARCH\FL_EUROPE_SCRAPE"
items = json.load(open(os.path.join(root, "state", "all_items.json"), encoding="utf-8"))

CMAP = {
    "uk": "UK", "united kingdom": "UK", "gb": "UK", "de": "Germany", "it": "Italy", "fr": "France",
    "es": "Spain", "nl": "Netherlands", "be": "Belgium", "hu": "Hungary", "jp": "Japan", "ro": "Romania",
    "czech republic": "Czechia",
}
def norm_country(c):
    c = (c or "").strip()
    if not c:
        return "Multi / n.a."
    k = c.lower()
    if k in CMAP:
        return CMAP[k]
    if any(s in c for s in [",", "/", ";", "+", "("]) or k.startswith(("multi", "pan", "12 countries", "none")):
        return "Multi-country"
    return c

THESIS = {"1_sale_exit": "Sale / exit", "2_closures": "Closures", "both": "Both", "against": "Against"}
REGION_APAC = {"Japan", "South Korea", "India", "Australia", "Australia/NZ", "Australia / New Zealand", "UAE", "AU/NZ", "JP"}

def clip(s, n):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"

rows = []
for d in items:
    ctry = norm_country(d.get("country"))
    region = "APAC" if (d.get("cluster") == "E" and d.get("avenue") != "licensing-template") or ctry in REGION_APAC or "Singapore" in str(d.get("country")) or "Thailand" in str(d.get("country")) or "Indonesia" in str(d.get("country")) else "Europe"
    rows.append({
        "id": d["id"],
        "tier": d.get("final_tier"),
        "status": d.get("status"),
        "thesis": THESIS.get(d.get("thesis"), d.get("thesis")),
        "region": region,
        "country": ctry,
        "date": clip(d.get("evidence_date"), 60),
        "claim": clip(d.get("claim"), 900),
        "data": clip(d.get("data_point"), 600),
        "quote": clip(d.get("english_translation") or d.get("original_quote"), 400),
        "source": clip(d.get("source_name"), 160),
        "url": str(d.get("source_url") or "").split(" ;")[0].split(" ")[0],
        "note": clip(d.get("verifier_note"), 420),
        "avenue": d.get("avenue"),
    })

json.dump(rows, open(os.path.join(root, "state", "items_for_page.json"), "w", encoding="utf-8"), ensure_ascii=False)
print("rows", len(rows))

# XLSX evidence log
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--user", "-q", "openpyxl"])
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

wb = Workbook()
cols = ["id", "tier", "status", "thesis", "region", "country", "date", "claim", "data", "quote", "source", "url", "note", "avenue"]
def sheet(name, rs):
    ws = wb.create_sheet(name)
    ws.append([c.upper() for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rs:
        ws.append([r[c] for c in cols])
    widths = {"id": 7, "tier": 5, "status": 11, "thesis": 11, "region": 8, "country": 14, "date": 18, "claim": 80, "data": 60, "quote": 50, "source": 30, "url": 40, "note": 50, "avenue": 22}
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths[c]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
wb.remove(wb.active)
sheet("Headline A-B", [r for r in rows if r["tier"] in ("A", "B") and r["thesis"] != "Against"])
sheet("Closures", [r for r in rows if r["thesis"] in ("Closures", "Both")])
sheet("Sale-exit", [r for r in rows if r["thesis"] in ("Sale / exit", "Both")])
sheet("Against", [r for r in rows if r["thesis"] == "Against"])
sheet("APAC", [r for r in rows if r["region"] == "APAC"])
sheet("All items", rows)
out = os.path.join(root, "FL_INTL_EVIDENCE_LOG.xlsx")
wb.save(out)
print("xlsx", out)

